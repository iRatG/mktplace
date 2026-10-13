"""Выгрузки для бухгалтерии в Excel (#59, #60): движения денег за период и реестр выплат.

По образцу `_legal_entities_xlsx` (apps/web/views/registration.py): строки в ячейках — как текст (защита от формул
из пользовательского ввода), даты — местное время без tzinfo. Итоги берутся из apps/billing/metrics.py.
"""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import metrics
from .models import Transaction

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

TOTAL_COLUMNS = [
    ("deposits", "Пополнения"),
    ("test_credits", "Тестовые начисления"),
    ("reserves", "Резервы"),
    ("releases", "Снятие резерва"),
    ("payments", "Оплаты по сделкам"),
    ("earnings", "Начисления исполнителям"),
    ("commission", "Комиссия платформы"),
    ("withdrawals", "Поставлено к выплате"),
    ("payouts", "Выплаты"),
    ("refunds", "Возвраты выплат"),
    ("corrections", "Корректировки"),
    ("tax", "Налог"),  # зарезервировано под НДФЛ (#47)
    ("count", "Операций"),
]

ROLE_LABELS = {"advertiser": "Рекламодатель", "blogger": "Исполнитель"}


def _naive(dt):
    return timezone.localtime(dt).replace(tzinfo=None) if dt else None


def _finish_sheet(ws, widths, text_rows_from=2, date_columns=()):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=text_rows_from):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for idx in date_columns:
            row[idx].number_format = "DD.MM.YYYY HH:MM"


def _save(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def finance_workbook(start, end):
    """Лист «Операции» — каждая транзакция за [start, end); лист «Итого» — по дням и общий итог. Возвращает bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Операции"
    ws.append(["Время", "Тип", "Пользователь", "Роль", "Сумма", "Баланс после", "Сделка", "Договор", "Оферта",
               "Описание"])
    for tx in metrics.transactions_between(start, end):
        user = tx.wallet.user
        ws.append([
            _naive(tx.created_at), tx.get_type_display(), user.login_display, ROLE_LABELS.get(user.role, user.role),
            tx.amount, tx.balance_after, tx.deal_id or "", tx.deal.contract_number if tx.deal_id else "",
            tx.offer_id or "", tx.description,
        ])
    _finish_sheet(ws, [18, 20, 34, 14, 16, 16, 10, 12, 10, 60], date_columns=(0,))

    ws2 = wb.create_sheet("Итого")
    ws2.append(["День"] + [label for _, label in TOTAL_COLUMNS])
    for day in metrics.daily_totals(start, end):
        ws2.append([day["day"].strftime("%d.%m.%Y")] + [day.get(key, "") for key, _ in TOTAL_COLUMNS])
    total = metrics.period_totals(start, end)
    ws2.append(["Итого"] + [total.get(key, "") for key, _ in TOTAL_COLUMNS])
    _finish_sheet(ws2, [12] + [18] * len(TOTAL_COLUMNS))
    return _save(wb)


def _requisites_text(snapshot):
    if not snapshot:
        return ""
    if snapshot.get("type") == "card":
        return f"Карта {snapshot.get('card_number', '')} {snapshot.get('card_holder_name', '')}".strip()
    return (f"Счёт {snapshot.get('account_number', '')} МФО {snapshot.get('bank_mfo', '')} "
            f"{snapshot.get('bank_name', '')} ИНН банка {snapshot.get('bank_inn', '')}").strip()


def payouts_workbook(payouts):
    """Реестр выплат: по строке на выплату с реквизитами и состоянием ЭСФ. Возвращает bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Выплаты"
    ws.append(["№", "Создана", "Вид", "Сделка", "Договор", "Исполнитель", "Категория", "Сумма", "Реквизиты",
               "ЭСФ", "Номер ЭСФ", "Блокировка", "Статус", "Проведена", "Комментарий / № платёжки"])
    for p in payouts:
        profile = getattr(p.blogger, "blogger_profile", None)
        einvoice = p.deal.performer_einvoice if p.deal_id else None
        ws.append([
            p.pk, _naive(p.created_at), p.get_kind_display(), p.deal_id or "",
            p.deal.contract_number if p.deal_id else "", p.blogger.login_display,
            profile.get_category_display() if profile else "", p.amount, _requisites_text(p.effective_requisites()),
            einvoice.get_status_display() if einvoice else "", einvoice.number if einvoice else "",
            ", ".join(p.blocked_reasons), p.get_status_display(), _naive(p.processed_at), p.admin_comment,
        ])
    _finish_sheet(ws, [8, 18, 18, 10, 12, 34, 18, 16, 50, 20, 18, 30, 14, 18, 40], date_columns=(1, 13))
    return _save(wb)


def xlsx_response(content, filename):
    from django.http import HttpResponse

    response = HttpResponse(content, content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


__all__ = ["finance_workbook", "payouts_workbook", "xlsx_response", "TOTAL_COLUMNS", "Transaction"]
