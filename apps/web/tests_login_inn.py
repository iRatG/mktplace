"""
Вход юрлица по ИНН (GitHub issue #2, openspec change legal-entity-login-by-inn).

Covers:
  - legal_entity_login / inn_from_login / parse_inn_login, User.login_display
  - вход по ИНН, по ИНН с пробелами, по старому служебному логину, по email
  - 8/10 цифр и неизвестный ИНН — общая ошибка, учёт в лимите login_fail
  - сотрудник видит ИНН как логин: выдача доступа, очередь, карточка, реестр, Excel
"""

import io

from django.contrib.auth.models import Group
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.registration.models import LegalEntityApplication
from apps.registration.services import (
    REGISTRATION_REVIEWERS_GROUP, inn_from_login, legal_entity_login, parse_inn_login,
)
from apps.users import security
from apps.users.models import User

INN = "123456789"
PASSWORD = "Ddocs1234!"


def _make_user(email, role=User.Role.ADVERTISER, is_staff=False):
    user = User.objects.create_user(email=email, password=PASSWORD, role=role)
    user.is_staff = is_staff
    user.is_email_confirmed = True
    user.status = User.Status.ACTIVE
    user.save(update_fields=["is_staff", "is_email_confirmed", "status"])
    return user


def _make_reviewer(email):
    reviewer = _make_user(email, is_staff=True)
    group, _ = Group.objects.get_or_create(name=REGISTRATION_REVIEWERS_GROUP)
    reviewer.groups.add(group)
    return reviewer


class LoginHelpersTest(TestCase):
    def test_login_roundtrip(self):
        self.assertEqual(legal_entity_login(INN), "legal.123456789@ddocs.internal")
        self.assertEqual(inn_from_login(legal_entity_login(INN)), INN)

    def test_inn_from_regular_email_is_none(self):
        self.assertIsNone(inn_from_login("user@demo.com"))
        self.assertIsNone(inn_from_login("legal.abc@ddocs.internal"))
        self.assertIsNone(inn_from_login(""))

    def test_parse_inn_login(self):
        self.assertEqual(parse_inn_login("123456789"), INN)
        self.assertEqual(parse_inn_login(" 123 456 789 "), INN)
        self.assertIsNone(parse_inn_login("12345678"))
        self.assertIsNone(parse_inn_login("1234567890"))
        self.assertIsNone(parse_inn_login("user@demo.com"))

    def test_login_display(self):
        self.assertEqual(_make_user(legal_entity_login(INN)).login_display, INN)
        self.assertEqual(_make_user("plain@demo.com").login_display, "plain@demo.com")


class LoginByInnTest(TestCase):
    def setUp(self):
        self.legal = _make_user(legal_entity_login(INN))
        self.url = reverse("web:login")

    def _login(self, login_value, password=PASSWORD):
        return self.client.post(self.url, {"email": login_value, "password": password})

    def _logged_in_as(self, user):
        return self.client.session.get("_auth_user_id") == str(user.pk)

    def test_login_page_field_is_text(self):
        html = self.client.get(self.url).content.decode()
        self.assertIn('type="text" name="email"', html)

    def test_login_by_inn(self):
        resp = self._login(INN)
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(self._logged_in_as(self.legal))

    def test_login_by_inn_with_spaces(self):
        self._login("123 456 789")
        self.assertTrue(self._logged_in_as(self.legal))

    def test_login_by_old_service_login(self):
        self._login("legal.123456789@ddocs.internal")
        self.assertTrue(self._logged_in_as(self.legal))

    def test_login_by_email_unchanged(self):
        regular = _make_user("Regular@Demo.com".lower())
        self._login("Regular@Demo.com")
        self.assertTrue(self._logged_in_as(regular))

    def test_wrong_password_generic_error(self):
        resp = self._login(INN, password="wrong")
        self.assertContains(resp, "Неверный логин или пароль.")
        self.assertFalse(self._logged_in_as(self.legal))

    def test_unknown_inn_generic_error(self):
        resp = self._login("987654321")
        self.assertContains(resp, "Неверный логин или пароль.")

    def test_eight_and_ten_digits_are_not_inn(self):
        for value in ("12345678", "1234567890"):
            resp = self._login(value)
            self.assertContains(resp, "Неверный логин или пароль.")
            self.assertFalse(self._logged_in_as(self.legal))


@override_settings(RATELIMIT_ENABLED=True)
class LoginByInnRateLimitTest(TestCase):
    def setUp(self):
        cache.clear()
        _make_user(legal_entity_login(INN))

    def tearDown(self):
        cache.clear()

    def test_failed_inn_attempts_count_in_ip_limit(self):
        url = reverse("web:login")
        for value, pw in (("987654321", PASSWORD), (INN, "wrong")):
            self.client.post(url, {"email": value, "password": pw}, REMOTE_ADDR="198.51.100.7")
        for _ in range(18):
            self.client.post(url, {"email": "987654321", "password": PASSWORD}, REMOTE_ADDR="198.51.100.7")
        self.assertTrue(security.is_limited("login_fail", "198.51.100.7", 20))


class StaffSeesInnAsLoginTest(TestCase):
    def setUp(self):
        self.reviewer = _make_reviewer("inn_reviewer@demo.com")
        self.application = LegalEntityApplication.objects.create(
            company_name="ООО Логин", inn=INN,
            status=LegalEntityApplication.Status.APPROVED, assigned_to=self.reviewer,
        )
        self.client.force_login(self.reviewer)

    def _issue(self):
        return self.client.post(
            reverse("web:admin_legal_entity_issue_access", args=[self.application.pk])
        )

    def test_issue_access_shows_inn_as_login(self):
        resp = self._issue()
        self.assertContains(resp, "Логин (ИНН)")
        self.assertContains(resp, INN)
        self.application.refresh_from_db()
        self.assertEqual(self.application.user.email, legal_entity_login(INN))

    def test_issued_credentials_work_with_inn(self):
        password = self._issue().context["issued_password"]
        self.client.logout()
        self.client.post(reverse("web:login"), {"email": INN, "password": password})
        self.application.refresh_from_db()
        self.assertEqual(self.client.session.get("_auth_user_id"), str(self.application.user_id))

    def test_detail_and_registry_show_inn_not_service_email(self):
        self._issue()
        detail = self.client.get(reverse("web:admin_legal_entity_detail", args=[self.application.pk]))
        registry = self.client.get(reverse("web:admin_legal_entity_registry"))
        for resp in (detail, registry):
            self.assertContains(resp, INN)
            self.assertNotContains(resp, "@ddocs.internal")

    def test_excel_export_login_is_inn(self):
        self._issue()
        resp = self.client.get(reverse("web:admin_legal_entity_registry") + "?export=xlsx")
        wb = load_workbook(io.BytesIO(resp.content))
        values = [str(c.value) for row in wb.active.iter_rows() for c in row if c.value is not None]
        self.assertIn(INN, values)
        self.assertFalse(any("@ddocs.internal" in v for v in values))
