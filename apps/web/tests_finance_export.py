"""
#59, #60, #61 (13.10.2026): страница «Финансы» и выгрузка в Excel; реестр выплат — фильтры, ЭСФ, проведение,
отказ, повторная выплата, остатки; страница сделки — договор и ЭСФ; плашка дашборда.
"""
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.billing.models import EInvoice, PayoutRequisites, Wallet, WithdrawalRequest
from apps.billing.services import BillingService
from apps.billing.tests_deal_payouts import DealPayoutBase, approved_card
from apps.profiles.models import BloggerProfile
from apps.users.models import PersonalDataAccess, User

D = Decimal


def _xlsx(response):
    return load_workbook(BytesIO(response.content))


class FinancePageTest(DealPayoutBase):
    def test_staff_only(self):
        self.client.force_login(self.blogger)
        self.assertEqual(self.client.get(reverse("web:admin_finance")).status_code, 302)
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(reverse("web:admin_finance")).status_code, 200)

    def test_totals_and_export(self):
        approved_card(self.blogger)
        deal = self._paid()
        BillingService.complete_withdrawal(WithdrawalRequest.objects.get(deal=deal), comment="п/п 5")
        self.client.force_login(self.staff)
        html = self.client.get(reverse("web:admin_finance")).content.decode()
        self.assertIn("Комиссия платформы", html)
        self.assertIn("16 000", html)
        self.assertIn("116 000", html)
        r = self.client.get(reverse("web:admin_finance") + "?export=xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("finance_", r["Content-Disposition"])
        wb = _xlsx(r)
        self.assertEqual(wb.sheetnames, ["Операции", "Итого"])
        ops = list(wb["Операции"].iter_rows(values_only=True))
        self.assertEqual(ops[0][:4], ("Время", "Тип", "Пользователь", "Роль"))
        self.assertGreaterEqual(len(ops), 5)
        self.assertIn(f"UB-{deal.pk}", [row[7] for row in ops[1:]])
        totals = list(wb["Итого"].iter_rows(values_only=True))
        self.assertEqual(totals[-1][0], "Итого")
        header = totals[0]
        self.assertEqual(totals[-1][header.index("Комиссия платформы")], D("16000"))
        self.assertEqual(totals[-1][header.index("Выплаты")], D("100000"))
        self.assertTrue(PersonalDataAccess.objects.filter(staff=self.staff, action__startswith="Выгрузка").exists())

    def test_period_swap_and_empty(self):
        self.client.force_login(self.staff)
        day = timezone.localdate()
        r = self.client.get(reverse("web:admin_finance"), {"date_from": day.isoformat(),
                                                           "date_to": (day - timedelta(days=2)).isoformat()})
        self.assertEqual(r.status_code, 200)
        self.assertEqual((r.context["date_from"], r.context["date_to"]), (day - timedelta(days=2), day))

    def test_dashboard_links(self):
        self.client.force_login(self.staff)
        html = self.client.get(reverse("web:admin_dashboard")).content.decode()
        self.assertIn("Выплат к проведению", html)
        self.assertIn(reverse("web:admin_finance"), html)


class PayoutRegisterTest(DealPayoutBase):
    def test_register_shows_blocked_and_ready(self):
        deal = self._paid()
        self.client.force_login(self.staff)
        html = self.client.get(reverse("web:admin_withdrawals")).content.decode()
        self.assertIn(f"UB-{deal.pk}", html)
        self.assertIn("нет подтверждённых реквизитов", html)
        self.assertIn("готовы к проведению: 0", html)
        approved_card(self.blogger)
        html = self.client.get(reverse("web:admin_withdrawals"), {"ready": "1"}).content.decode()
        self.assertIn("data-ready", html)
        self.assertTrue(PersonalDataAccess.objects.filter(staff=self.staff, action__startswith="Реестр").exists())

    def test_demo_hidden_by_default(self):
        demo = User.objects.create_user(email="demo@test.com", password="pass1234", role=User.Role.BLOGGER,
                                        status=User.Status.ACTIVE, is_demo=True)
        deal = self._paid(demo)
        self.client.force_login(self.staff)
        self.assertNotIn(f'data-payout-row="{deal.payout.pk}"',
                         self.client.get(reverse("web:admin_withdrawals")).content.decode())
        self.assertIn(f'data-payout-row="{deal.payout.pk}"',
                      self.client.get(reverse("web:admin_withdrawals"), {"demo": "1"}).content.decode())

    def test_approve_blocked_shows_reason_and_approve_ready_pays(self):
        deal = self._paid()
        payout = WithdrawalRequest.objects.get(deal=deal)
        self.client.force_login(self.staff)
        r = self.client.post(reverse("web:admin_withdrawal_approve", args=[payout.pk]),
                             {"comment": "п/п 1", "next": "/panel/withdrawals/?status=pending"}, follow=True)
        self.assertContains(r, "нет подтверждённых реквизитов")
        self.assertEqual(WithdrawalRequest.objects.get(pk=payout.pk).status, WithdrawalRequest.Status.PENDING)
        approved_card(self.blogger)
        r = self.client.post(reverse("web:admin_withdrawal_approve", args=[payout.pk]), {"comment": "п/п 1"},
                             follow=True)
        self.assertContains(r, "проведена")
        payout.refresh_from_db()
        self.assertEqual((payout.status, payout.admin_comment), (WithdrawalRequest.Status.COMPLETED, "п/п 1"))

    def test_einvoice_update_reject_and_reissue(self):
        BloggerProfile.objects.filter(user=self.blogger).update(category=BloggerProfile.Category.IP,
                                                                 is_ip_confirmed=True)
        deal = self._paid()
        einvoice = EInvoice.objects.get(deal=deal, kind="performer")
        self.client.force_login(self.staff)
        r = self.client.post(reverse("web:admin_einvoice_update", args=[einvoice.pk]),
                             {"status": "confirmed", "number": "ЭСФ-5", "next": f"/panel/withdrawals/?deal={deal.pk}"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(r["Location"], f"/panel/withdrawals/?deal={deal.pk}")
        einvoice.refresh_from_db()
        self.assertEqual((einvoice.status, einvoice.number), ("confirmed", "ЭСФ-5"))

        payout = WithdrawalRequest.objects.get(deal=deal)
        r = self.client.post(reverse("web:admin_withdrawal_reject", args=[payout.pk]), {"comment": ""}, follow=True)
        self.assertContains(r, "причину")
        self.client.post(reverse("web:admin_withdrawal_reject", args=[payout.pk]), {"comment": "Банк вернул"})
        self.assertEqual(WithdrawalRequest.objects.get(pk=payout.pk).status, WithdrawalRequest.Status.REJECTED)
        html = self.client.get(reverse("web:admin_withdrawals"), {"status": "rejected"}).content.decode()
        self.assertIn("Выплатить повторно", html)
        self.client.post(reverse("web:admin_withdrawal_reissue", args=[payout.pk]))
        self.assertEqual(deal.payouts.filter(status="pending").count(), 1)

    def test_legacy_balance_button(self):
        other = User.objects.create_user(email="old@test.com", password="pass1234", role=User.Role.BLOGGER,
                                         status=User.Status.ACTIVE)
        Wallet.objects.filter(user=other).update(available_balance=D("30000"))
        self.client.force_login(self.staff)
        html = self.client.get(reverse("web:admin_withdrawals")).content.decode()
        self.assertIn("data-legacy-balances", html)
        r = self.client.post(reverse("web:admin_withdrawal_legacy_balance", args=[other.pk]), follow=True)
        self.assertContains(r, "Сначала подтвердите реквизиты")
        approved_card(other)
        self.client.post(reverse("web:admin_withdrawal_legacy_balance", args=[other.pk]))
        wr = WithdrawalRequest.objects.get(blogger=other)
        self.assertEqual((wr.kind, wr.amount), (WithdrawalRequest.Kind.MANUAL, D("30000")))
        self.assertEqual(Wallet.objects.get(user=other).available_balance, D("0"))

    def test_platform_einvoice_queue_and_export(self):
        approved_card(self.blogger)
        deal = self._paid()
        BillingService.complete_withdrawal(WithdrawalRequest.objects.get(deal=deal), comment="п/п 8")
        self.client.force_login(self.staff)
        html = self.client.get(reverse("web:admin_withdrawals")).content.decode()
        self.assertIn("data-platform-einvoices", html)
        platform = EInvoice.objects.get(deal=deal, kind="platform")
        self.client.post(reverse("web:admin_einvoice_update", args=[platform.pk]),
                         {"status": "confirmed", "number": "П-1"})
        self.assertEqual(EInvoice.objects.get(pk=platform.pk).status, "confirmed")
        r = self.client.get(reverse("web:admin_withdrawals"), {"status": "all", "export": "xlsx"})
        rows = list(_xlsx(r)["Выплаты"].iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "№")
        self.assertEqual(rows[1][4], f"UB-{deal.pk}")
        self.assertEqual(rows[1][12], "Проведена")


class DealPageEsfTest(DealPayoutBase):
    def test_contract_and_esf_block_by_role(self):
        BloggerProfile.objects.filter(user=self.blogger).update(category=BloggerProfile.Category.IP,
                                                                 is_ip_confirmed=True)
        deal = self._paid()
        url = reverse("web:deal_detail", args=[deal.pk])
        self.client.force_login(self.blogger)
        html = self.client.get(url).content.decode()
        self.assertIn(f"договор № UB-{deal.pk}", html)
        self.assertIn("data-payout-status", html)
        self.assertIn("Выставьте рекламодателю ЭСФ", html)
        self.client.force_login(self.adv)
        html = self.client.get(url).content.decode()
        self.assertIn("Подтвердите ЭСФ исполнителя", html)
        self.client.force_login(self.staff)
        self.assertIn("Открыть в реестре выплат", self.client.get(url).content.decode())

    def test_individual_sees_not_required(self):
        deal = self._paid()
        self.client.force_login(self.blogger)
        html = self.client.get(reverse("web:deal_detail", args=[deal.pk])).content.decode()
        self.assertIn("ЭСФ не требуется", html)
        self.assertIn("ждёт: нет подтверждённых реквизитов", html)
